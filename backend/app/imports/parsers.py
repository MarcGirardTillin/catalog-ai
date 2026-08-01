"""Supplier-file parsers: raw bytes (PDF/XLSX/CSV) -> :class:`RawDocument`.

PDFs pass through untouched (Claude reads the pages directly). Tabular
files are parsed cell-by-cell so numeric values (EANs, prices) can be
cross-checked deterministically against the LLM extraction. Cell values
are preserved EXACTLY: an EAN read by Excel as the float 3607814866838.0
comes back as the string "3607814866838", never in scientific notation.
"""

import csv
import io
from decimal import Decimal
from pathlib import PurePath

import openpyxl  # type: ignore[import-untyped]

from app.imports.schema import RawDocument, RawLink, RawTable

_CSV_DELIMITERS = ",;\t"


def _cell_to_str(value: object) -> str:
    """Render one spreadsheet cell to its exact textual value."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, float):
        # Integer-valued floats (Excel stores big numbers like EANs as
        # floats): render without the trailing ".0" and never in
        # scientific notation.
        if value.is_integer():
            return str(int(value))
        text = repr(value)  # shortest round-trip representation
        if "e" in text or "E" in text:
            # Expand scientific notation into plain decimal form.
            text = format(Decimal(text), "f")
        return text
    return str(value)


def _parse_xlsx(data: bytes) -> list[RawTable]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        tables: list[RawTable] = []
        for sheet in workbook.worksheets:
            rows = [
                [_cell_to_str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            tables.append(RawTable(rows=rows, sheet=sheet.title))
        return tables
    finally:
        workbook.close()


def _decode_csv(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _parse_csv(data: bytes) -> RawTable:
    text = _decode_csv(data)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=_CSV_DELIMITERS)
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return RawTable(rows=[list(row) for row in reader], sheet=None)


def _pdf_links(data: bytes) -> list[RawLink]:
    """Liens cliquables (annotations URI) du PDF, avec leur page.

    Les bons de commande JOOR portent l'URL de la photo produit en bonne
    qualité derrière la vignette — invisible au rendu que lit Claude. Lecture
    best-effort : un PDF illisible par pypdf donne simplement zéro lien.
    """
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        links: list[RawLink] = []
        seen: set[tuple[int, str]] = set()
        for page_number, page in enumerate(reader.pages, start=1):
            for annotation in page.get("/Annots") or []:
                obj = annotation.get_object()
                action = obj.get("/A")
                uri = action.get("/URI") if action else None
                if not uri:
                    continue
                url = str(uri).strip()
                if not url.startswith(("http://", "https://")):
                    continue
                key = (page_number, url)
                if key in seen:
                    continue
                seen.add(key)
                links.append(RawLink(page=page_number, url=url))
        return links
    except Exception:  # noqa: BLE001 — bonus, jamais bloquant pour l'analyse
        return []


def parse_file(data: bytes, filename: str) -> RawDocument:
    """Parse a supplier file into a :class:`RawDocument`.

    Supported: ``.pdf`` (passthrough), ``.xlsx`` (openpyxl, all sheets),
    ``.csv`` (delimiter sniffed among ``,``/``;``/tab). ``.xls`` and any
    other extension raise :class:`ValueError`.
    """
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf":
        return RawDocument(
            kind="pdf", filename=filename, pdf_bytes=data, pdf_links=_pdf_links(data)
        )
    if suffix == ".xlsx":
        return RawDocument(kind="tabular", filename=filename, tables=_parse_xlsx(data))
    if suffix == ".csv":
        return RawDocument(kind="tabular", filename=filename, tables=[_parse_csv(data)])
    if suffix == ".xls":
        raise ValueError(
            "Legacy .xls files are not supported — convert the file to .xlsx first"
        )
    raise ValueError(
        f"Unsupported file type {suffix or '(none)'!r} for {filename!r} "
        "(expected .pdf, .xlsx or .csv)"
    )
