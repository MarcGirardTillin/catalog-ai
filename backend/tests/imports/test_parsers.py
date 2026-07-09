"""Tests for parse_file (synthetic fixtures built in memory)."""

import io

import openpyxl  # type: ignore[import-untyped]
import pytest

from app.imports.parsers import parse_file

VALID_EAN = "3607814866838"
OTHER_EAN = "4006381333931"


def _build_xlsx() -> bytes:
    """A small realistic purchase-order sheet."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Commande"
    sheet.append(
        ["Référence", "Désignation", "Coloris", "T36", "T38", "EAN", "PA HT", "PVP"]
    )
    # EAN written as an integer-valued float (how Excel usually stores it).
    sheet.append(["REF001", "Robe fleurie", "Rouge", 1, 2, float(VALID_EAN), 39.9, 89])
    sheet.append(["REF002", "Pantalon", None, None, 1, int(OTHER_EAN), 25.5, 59.95])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_xlsx_preserves_exact_values() -> None:
    document = parse_file(_build_xlsx(), "commande.xlsx")

    assert document.kind == "tabular"
    assert document.filename == "commande.xlsx"
    assert document.pdf_bytes is None
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.sheet == "Commande"
    assert table.rows[0][0] == "Référence"
    # Float-integer EAN comes back as the exact digit string.
    assert table.rows[1][5] == VALID_EAN
    assert table.rows[2][5] == OTHER_EAN
    # Decimals keep their value, integers have no trailing ".0".
    assert table.rows[1][6] == "39.9"
    assert table.rows[1][7] == "89"
    assert table.rows[2][7] == "59.95"
    # None cells become empty strings.
    assert table.rows[2][2] == ""
    assert table.rows[2][3] == ""


def test_parse_xlsx_extension_is_case_insensitive() -> None:
    document = parse_file(_build_xlsx(), "COMMANDE.XLSX")
    assert document.kind == "tabular"
    assert document.tables[0].rows[1][5] == VALID_EAN


def test_parse_csv_semicolon_delimiter_and_bom() -> None:
    csv_bytes = (
        "Référence;Désignation;EAN;PA HT;PVP\n"
        f"REF001;Robe fleurie;{VALID_EAN};39,90;89,00\n"
    ).encode("utf-8-sig")

    document = parse_file(csv_bytes, "commande.csv")

    assert document.kind == "tabular"
    assert len(document.tables) == 1
    rows = document.tables[0].rows
    assert rows[0] == ["Référence", "Désignation", "EAN", "PA HT", "PVP"]
    assert rows[1] == ["REF001", "Robe fleurie", VALID_EAN, "39,90", "89,00"]


def test_parse_csv_latin1_fallback() -> None:
    csv_bytes = "Réf,Désignation\nA1,Blouson été\n".encode("latin-1")
    document = parse_file(csv_bytes, "commande.csv")
    assert document.tables[0].rows[1] == ["A1", "Blouson été"]


def test_parse_pdf_is_passthrough() -> None:
    data = b"%PDF-1.4 fake supplier order"
    document = parse_file(data, "commande.pdf")
    assert document.kind == "pdf"
    assert document.pdf_bytes == data
    assert document.tables == []
    assert document.text is None


def test_parse_xls_raises_clear_error() -> None:
    with pytest.raises(ValueError, match="xlsx"):
        parse_file(b"\xd0\xcf\x11\xe0", "commande.xls")


def test_parse_unknown_extension_raises() -> None:
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_file(b"hello", "commande.docx")
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_file(b"hello", "noextension")
